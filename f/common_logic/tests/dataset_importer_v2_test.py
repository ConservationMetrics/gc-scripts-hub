import base64
import hashlib
import io
import json
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import fiona
import openpyxl
import psycopg
import pytest

import f.common_logic.dataset_importer_v2 as importer
from f.common_logic.dataset_importer_v2 import (
    ImportValidationError,
    apply_import,
    check_dataset_name,
    cleanup_expired_imports,
    list_datasets,
    preview_import,
    stage_import,
)
from f.common_logic.db_operations import StructuredDBWriter


def upload(name, text):
    return {"name": name, "data": base64.b64encode(text.encode()).decode()}


def upload_bytes(name, contents):
    return {"name": name, "data": base64.b64encode(contents).decode()}


def confirm(db, staged, preview):
    return apply_import(db, staged["import_id"], preview["preview_id"])


@pytest.fixture(autouse=True)
def importer_datalake(tmp_path, monkeypatch):
    datalake = tmp_path / "datalake"
    monkeypatch.setenv("DATASET_IMPORTER_DATALAKE_ROOT", str(datalake))
    return datalake


def table_rows(dsn, table):
    with psycopg.connect(dsn) as conn, conn.cursor() as cursor:
        cursor.execute(f'SELECT * FROM "public"."{table}" ORDER BY _id')
        columns = [column.name for column in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]


def test_create_csv_stages_previews_and_applies(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("birds.csv", "species,count\nheron,2\nibis,1\n"),
        "create",
        "Bird Observations",
    )

    assert staged["fields"] == ["species", "count"]
    assert staged["eligible_identity_fields"] == []
    assert check_dataset_name(mock_db_connection, "Bird Observations") == {
        "table_name": "bird_observations",
        "available": False,
    }
    preview = preview_import(mock_db_connection, staged["import_id"])
    assert {
        key: preview[key]
        for key in (
            "deleted",
            "updated",
            "added",
            "columns_added",
            "unchanged",
            "final_count",
        )
    } == {
        "deleted": 0,
        "updated": 0,
        "added": 2,
        "columns_added": 2,
        "unchanged": 0,
        "final_count": 2,
    }

    assert confirm(mock_db_connection, staged, preview)["success"] is True
    assert list_datasets(mock_db_connection) == ["bird_observations"]
    assert {
        tuple(sorted({key: row[key] for key in ("species", "count")}.items()))
        for row in table_rows(mock_db_connection, "bird_observations")
    } == {
        (("count", "2"), ("species", "heron")),
        (("count", "1"), ("species", "ibis")),
    }
    with pytest.raises(ImportValidationError, match="already been applied"):
        preview_import(mock_db_connection, staged["import_id"])


def test_csv_detects_gbif_tab_delimiter():
    rows = importer._parse_csv(
        b"gbifID\tscientificName\tlocality\n"
        b"1\tArdea alba\tNairobi, Kenya\n"
        b"2\tBubo africanus\tCape Town, South Africa\n"
    )

    assert rows == [
        {
            "gbifID": "1",
            "scientificName": "Ardea alba",
            "locality": "Nairobi, Kenya",
        },
        {
            "gbifID": "2",
            "scientificName": "Bubo africanus",
            "locality": "Cape Town, South Africa",
        },
    ]


def test_zip_detects_gbif_tab_delimited_csv():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "occurrence.csv",
            "gbifID\tscientificName\n1\tArdea alba\n2\tBubo africanus\n",
        )

    source_format, rows = importer._parse_zip(buffer.getvalue())

    assert source_format == "zip"
    assert rows == [
        {"gbifID": "1", "scientificName": "Ardea alba"},
        {"gbifID": "2", "scientificName": "Bubo africanus"},
    ]


def test_csv_still_rejects_inconsistent_row_width():
    with pytest.raises(
        ImportValidationError,
        match=r"^Invalid file\. CSV rows must match the header column count\.$",
    ):
        importer._parse_csv(b"species,count\nheron,2\nibis\n")


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_csv_preserves_bom_quoted_delimiters_and_multiline_cells(delimiter):
    contents = (
        f"\ufeffname{delimiter}note\r\n"
        f'Heron{delimiter}"comma, semicolon; tab\t and\r\nnewline"\r\n'
    ).encode()
    assert importer._parse_csv(contents) == [
        {"name": "Heron", "note": "comma, semicolon; tab\t and\r\nnewline"}
    ]


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
@pytest.mark.parametrize("values", ["Ibis", "Ibis{d}1{d}extra"])
def test_csv_fallback_rejects_inconsistent_rows(delimiter, values):
    contents = (
        f"name{delimiter}count\nHeron{delimiter}2\n{values.format(d=delimiter)}\n"
    )
    with pytest.raises(ImportValidationError, match="header column count"):
        importer._parse_csv(contents.encode())


@pytest.mark.parametrize("contents", [b"name,name\nHeron,2\n", b"name,\nHeron,2\n"])
def test_csv_rejects_invalid_headers(contents):
    with pytest.raises(ImportValidationError, match="column names"):
        importer._parse_csv(contents)


def test_csv_single_column_and_blank_lines():
    assert importer._parse_csv(b"name\n\nHeron\n") == [{"name": "Heron"}]
    assert importer._parse_csv(b"name\n") == []
    with pytest.raises(ImportValidationError, match="column names"):
        importer._parse_csv(b"")


def test_kobo_csv_stages_previews_and_applies(mock_db_connection, kobotoolbox_csv_file):
    staged = stage_import(
        mock_db_connection,
        upload_bytes(kobotoolbox_csv_file.name, kobotoolbox_csv_file.read_bytes()),
        "create",
        "kobo_rows",
    )
    assert staged["source_format"] == "csv"
    assert staged["record_count"] == 3
    preview = preview_import(mock_db_connection, staged["import_id"])
    assert confirm(mock_db_connection, staged, preview)["success"] is True
    rows = table_rows(mock_db_connection, "kobo_rows")
    assert any(
        "Arlington" in row.values() and "bamboo, wild boar" in row.values()
        for row in rows
    )


def test_zip_detects_kobo_semicolon_csv(kobotoolbox_csv_file):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(kobotoolbox_csv_file.name, kobotoolbox_csv_file.read_bytes())
    source_format, rows = importer._parse_zip(buffer.getvalue())
    assert source_format == "zip"
    assert len(rows) == 3
    assert any(
        "Arlington" in row.values() and "bamboo, wild boar" in row.values()
        for row in rows
    )


def test_csv_infers_gbif_point_geometry(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload(
            "occurrences.csv",
            "gbifID\tdecimalLongitude\tdecimalLatitude\tspecies\n"
            "1\t-54.137074\t3.248835\tAzeta mimica\n",
        ),
        "create",
        "gbif_points",
    )

    assert staged["fields"] == [
        "gbifID",
        "decimalLongitude",
        "decimalLatitude",
        "species",
        "g__type",
        "g__coordinates",
    ]
    preview = preview_import(mock_db_connection, staged["import_id"])
    assert preview["geometry_valid"] == 1
    assert preview["geometry_invalid"] == 0
    confirm(mock_db_connection, staged, preview)
    row = table_rows(mock_db_connection, "gbif_points")[0]
    assert row["g__type"] == "Point"
    assert row["g__coordinates"] == "[-54.137074,3.248835]"


@pytest.mark.parametrize(
    ("headers", "values"),
    [
        ("Longitude,Latitude", "-54,3"),
        ("lon,lat", "-54,3"),
        ("lng,lat", "-54,3"),
        ("Decimal Longitude,Decimal Latitude", "-54,3"),
    ],
)
def test_csv_recognizes_common_coordinate_aliases(headers, values):
    rows = importer._parse_csv(f"{headers}\n{values}\n".encode())

    assert rows[0]["g__type"] == "Point"
    assert rows[0]["g__coordinates"] == "[-54.0,3.0]"


def test_csv_sets_invalid_or_missing_coordinates_to_null(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload(
            "points.csv",
            "id,longitude,latitude\nvalid,-54,3\nmissing,,3\ninvalid,200,3\n",
        ),
        "create",
        "partial_points",
    )

    preview = preview_import(mock_db_connection, staged["import_id"])
    assert preview["geometry_valid"] == 1
    assert preview["geometry_invalid"] == 2
    confirm(mock_db_connection, staged, preview)
    rows = {row["id"]: row for row in table_rows(mock_db_connection, "partial_points")}
    assert (rows["valid"]["g__type"], rows["valid"]["g__coordinates"]) == (
        "Point",
        "[-54.0,3.0]",
    )
    assert (rows["missing"]["g__type"], rows["missing"]["g__coordinates"]) == (
        None,
        None,
    )
    assert (rows["invalid"]["g__type"], rows["invalid"]["g__coordinates"]) == (
        None,
        None,
    )


def test_csv_rejects_ambiguous_coordinate_pairs():
    with pytest.raises(ImportValidationError, match="Multiple longitude"):
        importer._parse_csv(b"longitude,latitude,lon,lat\n-54,3,-54,3\n")


def test_csv_warns_when_a_coordinate_pair_is_incomplete(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("points.csv", "id,longitude\nA,-54\n"),
        "create",
        "incomplete_points",
    )

    assert "no map geometry was generated" in staged["geometry_warning"]
    assert "g__type" not in staged["fields"]


def test_csv_preserves_valid_explicit_geometry():
    rows = importer._parse_csv(
        b'id,g__type,g__coordinates\nA,Point,"[\"\"-54\"\",\"\"3\"\"]"\nB,,\n'
    )

    assert rows[0]["g__coordinates"] == "[-54.0,3.0]"
    assert rows[1]["g__type"] is None
    assert rows[1]["g__coordinates"] is None


def test_csv_rejects_partial_explicit_geometry_columns():
    with pytest.raises(ImportValidationError, match="include both"):
        importer._parse_csv(b"id,g__type\nA,Point\n")


def test_csv_rejects_malformed_explicit_non_point_geometry():
    with pytest.raises(ImportValidationError, match="invalid coordinates"):
        importer._parse_csv(
            b'id,g__type,g__coordinates\nA,LineString,"[1,2]"\n'
        )


def test_merge_preserves_omitted_values_and_imported_null_clears_them(
    mock_db_connection,
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT, note TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s, %s)',
            ("one", "A", "old"),
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s, %s)',
            ("two", "B", "retain"),
        )

    staged = stage_import(
        mock_db_connection,
        upload("update.csv", "code,note\nA,\nC,new\n"),
        "merge",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert preview["added"] == 1
    assert preview["updated"] == 1
    assert preview["unchanged"] == 1
    confirm(mock_db_connection, staged, preview)
    rows = table_rows(mock_db_connection, "observations")
    assert {(row["code"], row["note"]) for row in rows} == {
        ("A", None),
        ("B", "retain"),
        ("C", "new"),
    }


@pytest.mark.parametrize("source_format", ["shapefile", "geopackage"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
def test_spatial_null_attributes_follow_update_policy(
    mock_db_connection, tmp_path, source_format, policy
):
    from f.common_logic.data_conversion import convert_data

    path = tmp_path / ("birds.shp" if source_format == "shapefile" else "birds.gpkg")
    with fiona.open(
        path,
        "w",
        driver="ESRI Shapefile" if source_format == "shapefile" else "GPKG",
        schema={"geometry": "Point", "properties": {"code": "str", "note": "str"}},
        crs="EPSG:4326",
    ) as dataset:
        dataset.write(
            {
                "geometry": {"type": "Point", "coordinates": [1, 2]},
                "properties": {"code": "A", "note": None},
            }
        )
    # Existing conversion callers still omit NULL attributes by default.
    converted, _ = convert_data([str(path)], source_format)
    assert "note" not in converted["features"][0]["properties"]
    if source_format == "shapefile":
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            for member in tmp_path.glob("birds.*"):
                archive.write(member, member.name)
        payload = upload_bytes("birds.zip", buffer.getvalue())
    else:
        payload = upload_bytes(path.name, path.read_bytes())

    with psycopg.connect(mock_db_connection) as conn:
        conn.execute(
            """CREATE TABLE observations (
                _id TEXT PRIMARY KEY, code TEXT, note TEXT, feature_id TEXT,
                g__type TEXT, g__coordinates TEXT, __geopackage_layer TEXT
            )"""
        )
        conn.execute(
            """INSERT INTO observations VALUES
                ('one', 'A', 'old note', '1', 'Point', '[1.0,2.0]', 'birds')"""
        )
    staged = stage_import(mock_db_connection, payload, "merge", "observations")
    assert "note" in staged["fields"]
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"], policy)
    assert preview["updated"] == (1 if policy == "imported" else 0)
    assert preview["unchanged"] == (0 if policy == "imported" else 1)
    assert preview["added"] == preview["columns_added"] == 0
    confirm(mock_db_connection, staged, preview)
    assert table_rows(mock_db_connection, "observations")[0]["note"] == (
        None if policy == "imported" else "old note"
    )


def test_existing_wins_preserves_matches_and_adds_new_rows(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT, note TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s, %s), (%s, %s, %s)',
            ("one", "A", "old", "two", "B", "retain"),
        )

    staged = stage_import(
        mock_db_connection,
        upload("update.csv", "code,note\nA,new\nC,added\n"),
        "merge",
        "observations",
    )
    preview = preview_import(
        mock_db_connection, staged["import_id"], ["code"], "existing"
    )
    assert preview["added"] == 1
    assert preview["updated"] == 0
    assert preview["unchanged"] == 2
    confirm(mock_db_connection, staged, preview)

    assert {(row["code"], row["note"]) for row in table_rows(mock_db_connection, "observations")} == {
        ("A", "old"),
        ("B", "retain"),
        ("C", "added"),
    }


def test_duplicate_identity_is_rejected_before_writes(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
    staged = stage_import(
        mock_db_connection,
        upload("duplicate.csv", "code\nA\nA\n"),
        "merge",
        "observations",
    )
    with pytest.raises(
        ImportValidationError, match="Uploaded identity combinations must be unique"
    ):
        preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert table_rows(mock_db_connection, "observations") == []


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
def test_duplicate_composite_identity_in_target_is_rejected(
    mock_db_connection, goal, policy
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."people" (_id TEXT PRIMARY KEY, first_name TEXT, last_name TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."people" VALUES (%s, %s, %s), (%s, %s, %s)',
            ("one", "Ada", "Lovelace", "two", "Ada", "Lovelace"),
        )
    staged = stage_import(
        mock_db_connection,
        upload("people.csv", "first_name,last_name\nAda,Lovelace\n"),
        goal,
        "people",
    )

    with pytest.raises(ImportValidationError, match="target dataset"):
        preview_import(
            mock_db_connection,
            staged["import_id"],
            ["first_name", "last_name"],
            policy,
        )

    assert len(table_rows(mock_db_connection, "people")) == 2


@pytest.mark.parametrize(
    "identity_fields",
    [
        ["first_name", "first_name"],
        ["first_name", "last_name", "region", "year"],
    ],
)
def test_identity_requires_one_to_three_distinct_fields(
    mock_db_connection, identity_fields
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."people" (_id TEXT PRIMARY KEY, first_name TEXT, last_name TEXT, region TEXT, year TEXT)'
        )
    staged = stage_import(
        mock_db_connection,
        upload(
            "people.csv",
            "first_name,last_name,region,year\nAda,Lovelace,London,1843\n",
        ),
        "merge",
        "people",
    )

    with pytest.raises(ImportValidationError, match="one to three distinct"):
        preview_import(mock_db_connection, staged["import_id"], identity_fields)


def test_geojson_preserves_geometry_and_nested_properties(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload(
            "birds.geojson",
            '{"type":"FeatureCollection","features":[{"type":"Feature","id":"feature-123","properties":{"name":"heron","tags":["wetland"]},"geometry":{"type":"Point","coordinates":[1,2]}}]}',
        ),
        "create",
        "birds",
    )
    assert staged["fields"] == [
        "name",
        "tags",
        "feature.id",
        "g__type",
        "g__coordinates",
    ]
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)
    row = table_rows(mock_db_connection, "birds")[0]
    assert row["g__type"] == "Point"
    assert row["g__coordinates"] == "[1,2]"
    assert row["feature_id"] == "feature-123"
    assert row["tags"] == '["wetland"]'


@pytest.mark.parametrize("filename", ["mixed.geojson", "mixed.json"])
def test_geojson_rejects_geometry_collections_during_staging(
    mock_db_connection, filename
):
    source = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "properties": {"name": "mixed"},
                "geometry": {
                    "type": "GeometryCollection",
                    "geometries": [
                        {"type": "Point", "coordinates": [1, 2]},
                        {
                            "type": "LineString",
                            "coordinates": [[3, 4], [5, 6]],
                        },
                    ],
                },
            }
        ],
    }

    with pytest.raises(
        ImportValidationError,
        match="GeometryCollection geometries are not supported by this importer",
    ):
        stage_import(
            mock_db_connection,
            upload(filename, json.dumps(source)),
            "create",
            "mixed_geometries",
        )
    assert check_dataset_name(mock_db_connection, "mixed_geometries")["available"] is True


@pytest.mark.parametrize("goal", ["merge", "sync"])
def test_confirmation_rejects_a_target_changed_after_review(mock_db_connection, goal):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s)', ("one", "A")
        )
    staged = stage_import(
        mock_db_connection,
        upload("update.csv", "code\nB\n"),
        goal,
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s)', ("two", "C")
        )
    with pytest.raises(ImportValidationError, match="changed after review"):
        confirm(mock_db_connection, staged, preview)
    assert {
        (row["_id"], row["code"])
        for row in table_rows(mock_db_connection, "observations")
    } == {
        ("one", "A"),
        ("two", "C"),
    }


def test_only_compatible_datasets_are_listed_and_targetable(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
        cursor.execute('CREATE TABLE "public"."application_state" (name TEXT)')
        cursor.execute(
            'CREATE TABLE "public"."observations__metadata" (_id TEXT PRIMARY KEY)'
        )

    assert list_datasets(mock_db_connection) == ["observations"]
    with pytest.raises(ImportValidationError, match="not a compatible"):
        stage_import(
            mock_db_connection,
            upload("rows.csv", "name\nunsafe\n"),
            "append",
            "application_state",
        )


def test_create_persists_column_mapping_for_later_imports(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("birds.csv", "Bird Name,Count\nHeron,2\n"),
        "create",
        "Bird Observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'SELECT original_column, sql_column FROM "public"."bird_observations__columns" ORDER BY original_column'
        )
        assert cursor.fetchall() == [("Bird Name", "Bird_Name"), ("Count", "Count")]

    appended = stage_import(
        mock_db_connection,
        upload("more.csv", "Bird Name,Count\nIbis,1\n"),
        "append",
        "bird_observations",
    )
    appended_preview = preview_import(mock_db_connection, appended["import_id"])
    confirm(mock_db_connection, appended, appended_preview)
    assert {
        (row["Bird_Name"], row["Count"])
        for row in table_rows(mock_db_connection, "bird_observations")
    } == {
        ("Heron", "2"),
        ("Ibis", "1"),
    }


@pytest.mark.parametrize("goal", ["merge", "sync"])
def test_staging_only_offers_existing_columns_for_identity(mock_db_connection, goal):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE public.observations (_id TEXT PRIMARY KEY, code TEXT)'
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,new_field,_id\nA,new,external\n"),
        goal,
        "observations",
    )
    assert staged["fields"] == ["code", "new_field", "_id"]
    assert staged["eligible_identity_fields"] == ["code"]
    for field in ("new_field", "_id"):
        with pytest.raises(ImportValidationError, match="do not exist in the target"):
            preview_import(mock_db_connection, staged["import_id"], [field])
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert preview["added"] == 1
    assert preview["columns_added"] == 2
    assert table_rows(mock_db_connection, "observations") == []


@pytest.mark.parametrize("goal", ["merge", "sync"])
def test_staging_succeeds_without_shared_identity_columns(mock_db_connection, goal):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute('CREATE TABLE public.observations (_id TEXT PRIMARY KEY)')
    staged = stage_import(
        mock_db_connection, upload("rows.csv", "code\nA\n"), goal, "observations"
    )
    assert staged["record_count"] == 1
    assert staged["eligible_identity_fields"] == []


def test_existing_column_mapping_is_used_for_identity(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, bird_name TEXT, note TEXT)'
        )
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column VARCHAR(128), sql_column VARCHAR(64) NOT NULL)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations__columns" VALUES (%s, %s)',
            ("Bird Name", "bird_name"),
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s, %s)',
            ("one", "Heron", "old"),
        )

    staged = stage_import(
        mock_db_connection,
        upload("birds.csv", "Bird Name,note\nHeron,new\n"),
        "merge",
        "observations",
    )
    assert staged["eligible_identity_fields"] == ["Bird Name", "note"]
    preview = preview_import(
        mock_db_connection, staged["import_id"], ["Bird Name"]
    )
    assert preview["updated"] == 1
    confirm(mock_db_connection, staged, preview)
    assert table_rows(mock_db_connection, "observations")[0]["note"] == "new"


def test_legacy_mapping_table_is_upgraded_for_long_source_columns(
    mock_db_connection,
):
    long_column = "field_" + ("x" * 140)
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY)'
        )
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column VARCHAR(128), sql_column VARCHAR(64) NOT NULL)'
        )

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", f"{long_column}\nvalue\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            """SELECT data_type FROM information_schema.columns
               WHERE table_schema = 'public'
                 AND table_name = 'observations__columns'
                 AND column_name = 'original_column'"""
        )
        assert cursor.fetchone() == ("text",)
        cursor.execute(
            'SELECT sql_column FROM "public"."observations__columns" WHERE original_column = %s',
            (long_column,),
        )
        assert cursor.fetchone() is not None


def test_legacy_internal_id_mapping_does_not_capture_uploaded_id(
    mock_db_connection,
):
    writer = StructuredDBWriter(
        mock_db_connection, "observations", use_mapping_table=True
    )
    writer.handle_output([{"_id": "legacy", "name": "Heron"}])

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "_id,name\nexternal,Ibis\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    rows = table_rows(mock_db_connection, "observations")
    imported = next(row for row in rows if row["name"] == "Ibis")
    assert imported["_id"] != "external"
    assert imported["source_id"] == "external"

    writer.handle_output([{"_id": "legacy", "name": "updated"}])
    assert next(row for row in table_rows(mock_db_connection, "observations") if row["_id"] == "legacy")["name"] == "updated"


def test_v2_reuses_unmapped_columns_from_default_legacy_writer(mock_db_connection):
    writer = StructuredDBWriter(mock_db_connection, "observations")
    writer.handle_output([{"_id": "legacy", "Bird Name": "Heron"}])

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "Bird Name\nIbis\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_schema = 'public' AND table_name = 'observations'"
        )
        columns = {row[0] for row in cursor.fetchall()}
    assert "BirdName" in columns
    assert "Bird_Name" not in columns


def test_unmapped_fallback_does_not_claim_a_persisted_mapping(
    mock_db_connection,
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, "BirdName" TEXT)'
        )
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column TEXT, sql_column TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations__columns" VALUES (%s, %s)',
            ("Old Name", "BirdName"),
        )

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "Bird Name\nIbis\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    row = table_rows(mock_db_connection, "observations")[0]
    assert row["BirdName"] is None
    assert row["Bird_Name"] == "Ibis"


def test_legacy_writer_and_v2_round_trip_through_shared_mappings(
    mock_db_connection,
):
    writer = StructuredDBWriter(
        mock_db_connection,
        "observations",
        use_mapping_table=True,
        reverse_properties_separated_by="/",
        sep_policy="underscore",
    )
    writer.handle_output(
        [
            {
                "_id": "one",
                "Basic information/Bird Name": "Heron",
                "Bird-Name": "old",
            }
        ]
    )

    staged = stage_import(
        mock_db_connection,
        upload(
            "birds.csv",
            "Basic information/Bird Name,Bird-Name\nHeron,from-v2\n",
        ),
        "merge",
        "observations",
    )
    preview = preview_import(
        mock_db_connection,
        staged["import_id"],
        ["Basic information/Bird Name"],
    )
    confirm(mock_db_connection, staged, preview)

    writer.handle_output(
        [
            {
                "_id": "one",
                "Basic information/Bird Name": "Heron",
                "Bird-Name": "from-legacy",
            }
        ]
    )

    rows = table_rows(mock_db_connection, "observations")
    assert len(rows) == 1
    assert rows[0]["Bird_Name__Basic_information"] == "Heron"
    assert rows[0]["Bird_Name"] == "from-legacy"


@pytest.mark.parametrize("dataset_name", ["", "   ", "---"])
def test_dataset_name_requires_letters_or_numbers(mock_db_connection, dataset_name):
    with pytest.raises(ImportValidationError, match="dataset name"):
        check_dataset_name(mock_db_connection, dataset_name)


def test_stage_rejects_invalid_and_reserved_create_names(mock_db_connection):
    for name in ("---", "observations__columns"):
        with pytest.raises(ImportValidationError, match="dataset name|reserved suffix"):
            stage_import(
                mock_db_connection,
                upload("rows.csv", "name\nHeron\n"),
                "create",
                name,
            )


def test_dataset_name_is_unavailable_when_any_relation_uses_it(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute('CREATE VIEW "public"."observations" AS SELECT 1 AS value')

    assert check_dataset_name(mock_db_connection, "observations")["available"] is False


def test_persisted_mapping_is_not_overwritten_by_fallback(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT, code_raw TEXT)'
        )
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column TEXT, sql_column TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations__columns" VALUES (%s, %s)',
            ("code", "code_raw"),
        )

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code\nA\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)
    row = table_rows(mock_db_connection, "observations")[0]
    assert row["code"] is None
    assert row["code_raw"] == "A"


def test_expired_sessions_are_rejected_and_cleaned_up(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET expires_at = now() - INTERVAL '1 second' WHERE import_id = %s",
            (staged["import_id"],),
        )

    assert cleanup_expired_imports(mock_db_connection) == 1
    with pytest.raises(ImportValidationError, match="expired"):
        preview_import(mock_db_connection, staged["import_id"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT 1 FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        assert cursor.fetchone() is None
    assert check_dataset_name(mock_db_connection, "observations")["available"] is True


def test_previous_importer_schema_is_adopted_with_staged_sessions(
    mock_db_connection, monkeypatch
):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "ALTER SCHEMA dataset_importer_v2 RENAME TO dataset_importer_old"
        )
        assert importer._schema_has_importer_contract(
            cursor, "dataset_importer_old"
        )
        cursor.execute(
            "SELECT pg_get_userbyid(nspowner), current_user FROM pg_namespace WHERE nspname = 'dataset_importer_old'"
        )
        owner, current_user = cursor.fetchone()
        assert owner == current_user
    monkeypatch.setenv(
        "DATASET_IMPORTER_V2_MIGRATE_SCHEMA", "dataset_importer_old"
    )

    with ThreadPoolExecutor(max_workers=2) as executor:
        listings = list(
            executor.map(lambda _: list_datasets(mock_db_connection), range(2))
        )
    assert listings == [[], []]

    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("SELECT to_regnamespace('dataset_importer_v2') IS NOT NULL")
        assert cursor.fetchone() == (True,)
        cursor.execute("SELECT to_regnamespace('dataset_importer_old')")
        assert cursor.fetchone() == (None,)
    assert table_rows(mock_db_connection, "observations")[0]["name"] == "Heron"


def test_unrelated_schema_is_not_adopted(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("CREATE SCHEMA dataset_importer_bad")
        cursor.execute("CREATE TABLE dataset_importer_bad.import_sessions (id TEXT)")
        cursor.execute("CREATE TABLE dataset_importer_bad.import_rows (id TEXT)")
        cursor.execute("CREATE TABLE dataset_importer_bad.dataset_registry (id TEXT)")

    assert list_datasets(mock_db_connection) == []

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("SELECT to_regnamespace('dataset_importer_bad') IS NOT NULL")
        assert cursor.fetchone() == (True,)
        cursor.execute("SELECT to_regnamespace('dataset_importer_v2') IS NOT NULL")
        assert cursor.fetchone() == (True,)


def test_dataset_listing_cleans_expired_sources_on_app_load(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET expires_at = now() - INTERVAL '1 second' WHERE import_id = %s",
            (staged["import_id"],),
        )

    assert list_datasets(mock_db_connection) == []
    assert check_dataset_name(mock_db_connection, "observations")["available"] is True


def test_expiry_cleanup_removes_an_archive_orphaned_before_commit(
    mock_db_connection, importer_datalake
):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview_import(mock_db_connection, staged["import_id"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT target_table, source_name, source_data FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        table_name, source_name, source_data = cursor.fetchone()
        archive_path, created = importer._write_source_archive(
            table_name, source_name, source_data, staged["import_id"]
        )
        assert created is True
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET expires_at = now() - INTERVAL '1 second' WHERE import_id = %s",
            (staged["import_id"],),
        )

    assert Path(archive_path).is_relative_to(importer_datalake)
    assert Path(archive_path).exists()
    assert cleanup_expired_imports(mock_db_connection) == 1
    assert not Path(archive_path).exists()


def test_expiry_cleanup_preserves_successful_archives(
    mock_db_connection, importer_datalake
):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    archive_path = Path(confirm(mock_db_connection, staged, preview)["archive_path"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET expires_at = now() - INTERVAL '1 second' WHERE import_id = %s",
            (staged["import_id"],),
        )

    assert archive_path.is_relative_to(importer_datalake)
    assert cleanup_expired_imports(mock_db_connection) == 1
    assert archive_path.read_bytes() == b"name\nHeron\n"


def test_expiry_cleanup_retains_session_when_orphan_removal_fails(
    mock_db_connection, monkeypatch
):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET expires_at = now() - INTERVAL '1 second' WHERE import_id = %s",
            (staged["import_id"],),
        )

    def fail_removal(_path):
        raise OSError("storage unavailable")

    monkeypatch.setattr(importer, "_remove_archive_file", fail_removal)
    assert cleanup_expired_imports(mock_db_connection) == 0
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT source_data IS NOT NULL FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        assert cursor.fetchone() == (True,)


def test_mapping_change_after_staging_rejects_apply(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT, alternate TEXT)'
        )

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code\nA\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column TEXT, sql_column TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations__columns" VALUES (%s, %s)',
            ("code", "alternate"),
        )

    with pytest.raises(ImportValidationError, match="mapping changed"):
        confirm(mock_db_connection, staged, preview)


def test_create_rejects_an_orphaned_mapping_table(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations__columns" (original_column TEXT, sql_column TEXT)'
        )

    assert check_dataset_name(mock_db_connection, "observations")["available"] is False
    with pytest.raises(ImportValidationError, match="mapping|conflicts"):
        stage_import(
            mock_db_connection,
            upload("rows.csv", "name\nHeron\n"),
            "create",
            "observations",
        )


def test_staged_create_reserves_its_mapping_table_name(mock_db_connection):
    shared_prefix = "a" * 54
    first_name = f"{shared_prefix}_first"
    second_name = f"{shared_prefix}_second"
    stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        first_name,
    )

    with pytest.raises(ImportValidationError, match="conflicts"):
        stage_import(
            mock_db_connection,
            upload("rows.csv", "name\nIbis\n"),
            "create",
            second_name,
        )


def test_existing_long_dataset_reserves_its_mapping_name(mock_db_connection):
    shared_prefix = "a" * 54
    existing_name = f"{shared_prefix}_existing"
    new_name = f"{shared_prefix}_new"
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            psycopg.sql.SQL("CREATE TABLE {} (_id TEXT PRIMARY KEY)").format(
                psycopg.sql.Identifier(existing_name)
            )
        )

    assert check_dataset_name(mock_db_connection, new_name)["available"] is False


def test_existing_staged_create_is_backfilled_into_registry(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "DELETE FROM dataset_importer_v2.dataset_registry WHERE target_table = 'observations'"
        )

    preview_import(mock_db_connection, staged["import_id"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT status, source_import_id FROM dataset_importer_v2.dataset_registry WHERE target_table = 'observations'"
        )
        status, source_import_id = cursor.fetchone()
        assert status == "reserved"
        assert str(source_import_id) == staged["import_id"]


@pytest.mark.parametrize(
    ("relative_path", "expected_format"),
    [
        ("locusmap_favorites.gpx", "gpx"),
        ("locusmap_favorites.kml", "kml"),
        ("smart_patrol_sample.xml", "smart"),
        ("my_shapefile_data.zip", "shapefile"),
    ],
)
def test_supported_formats_are_staged(
    mock_db_connection, relative_path, expected_format
):
    path = Path(__file__).parent / "assets" / relative_path
    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        "create",
        f"dataset_{expected_format}",
    )
    assert staged["source_format"] == expected_format
    assert staged["record_count"] > 0


def test_single_layer_geopackage_is_staged(mock_db_connection, tmp_path):
    path = tmp_path / "points.gpkg"
    with fiona.open(
        path,
        "w",
        driver="GPKG",
        layer="points",
        schema={"geometry": "Point", "properties": {"name": "str"}},
        crs="EPSG:4326",
    ) as collection:
        collection.write(
            {
                "geometry": {"type": "Point", "coordinates": (1, 2)},
                "properties": {"name": "Heron"},
            }
        )
    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        "create",
        "geopackage_rows",
    )
    assert staged["source_format"] == "geopackage"
    assert staged["record_count"] == 1


def test_multi_layer_geopackage_is_imported(mock_db_connection):
    path = Path(__file__).parent / "assets" / "datasets_bees.gpkg"
    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        "create",
        "geopackage_rows",
    )
    assert staged["source_format"] == "geopackage"
    assert staged["record_count"] == 55

    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)
    rows = table_rows(mock_db_connection, "geopackage_rows")
    assert len(rows) == 55
    for layer, geometry_type, count in (
        ("apiary", "Point", 36),
        ("area", "Polygon", 18),
        ("tracks", "LineString", 1),
    ):
        layer_rows = [row for row in rows if row["__geopackage_layer"] == layer]
        assert len(layer_rows) == count
        assert all(row["g__type"] == geometry_type for row in layer_rows)
    assert {row["__geopackage_layer"] for row in rows} == {
        "apiary",
        "area",
        "tracks",
    }


def test_single_sheet_xlsx_is_staged(mock_db_connection, tmp_path):
    path = tmp_path / "birds.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["name", "count"])
    sheet.append(["Heron", 2])
    workbook.save(path)
    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        "create",
        "xlsx_rows",
    )
    assert staged["source_format"] == "xlsx"
    assert staged["record_count"] == 1


def test_xlsx_infers_point_geometry(mock_db_connection, tmp_path):
    path = tmp_path / "points.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["name", "longitude", "latitude"])
    sheet.append(["Heron", -54.1, 3.2])
    workbook.save(path)

    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        "create",
        "xlsx_points",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])

    assert preview["geometry_valid"] == 1
    assert preview["geometry_invalid"] == 0


@pytest.mark.parametrize(
    ("goal", "preview_error"),
    [("append", None), ("merge", None), ("sync", "cannot be empty")],
)
def test_empty_xlsx_uses_operation_empty_rules(
    mock_db_connection, tmp_path, goal, preview_error
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s)', ("one", "A")
        )
    path = tmp_path / "empty.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["code"])
    workbook.save(path)
    staged = stage_import(
        mock_db_connection,
        upload_bytes(path.name, path.read_bytes()),
        goal,
        "observations",
    )
    assert staged["record_count"] == 0

    if preview_error:
        with pytest.raises(ImportValidationError, match=preview_error):
            preview_import(mock_db_connection, staged["import_id"])
    else:
        preview = preview_import(mock_db_connection, staged["import_id"])
        assert preview["added"] == preview["updated"] == preview["deleted"] == 0
        confirm(mock_db_connection, staged, preview)
        assert table_rows(mock_db_connection, "observations")[0]["code"] == "A"


def test_json_and_cybertracker_are_detected_by_content(mock_db_connection):
    regular = stage_import(
        mock_db_connection,
        upload("rows.json", '[{"name":"Heron","tags":["wetland"]}]'),
        "create",
        "json_rows",
    )
    assert regular["source_format"] == "json"
    cybertracker = (
        Path(__file__).resolve().parents[2]
        / "connectors"
        / "cybertracker"
        / "tests"
        / "assets"
        / "0.json"
    )
    detected = stage_import(
        mock_db_connection,
        upload_bytes(cybertracker.name, cybertracker.read_bytes()),
        "create",
        "cybertracker_rows",
    )
    assert detected["source_format"] == "cybertracker"
    assert detected["record_count"] > 0


@pytest.mark.parametrize("document", ["[]", "{}", '[{"name":"Heron"}, 1]'])
def test_json_requires_a_nonempty_array_of_objects(mock_db_connection, document):
    with pytest.raises(ImportValidationError, match="array of objects|record must be an object"):
        stage_import(
            mock_db_connection,
            upload("rows.json", document),
            "create",
            "json_rows",
        )


def test_unsupported_top_level_file_is_rejected(mock_db_connection):
    with pytest.raises(ImportValidationError, match="not supported"):
        stage_import(
            mock_db_connection,
            upload("rows.txt", "name\nHeron\n"),
            "create",
            "observations",
        )


def test_sanitized_column_name_collisions_use_legacy_suffixes(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "Bird Name,Bird-Name\nHeron,Ibis\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    row = table_rows(mock_db_connection, "observations")[0]
    assert row["Bird_Name"] == "Heron"
    assert row["Bird_Name_001"] == "Ibis"


def test_create_uses_legacy_nested_and_metadata_column_names(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload(
            "rows.csv",
            "Basic information/Are you married?,$categoryId,categoryId\nyes,system,user\n",
        ),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    row = table_rows(mock_db_connection, "observations")[0]
    assert row["Are_you_married__Basic_information"] == "yes"
    assert row["__categoryId"] == "system"
    assert row["categoryId"] == "user"


def test_uploaded_id_and_source_id_are_kept_as_distinct_data(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "_id,source_id\nexternal-id,source-value\n"),
        "create",
        "observations",
    )
    assert staged["source_mapping"] == {
        "_id": "source_id",
        "source_id": "source_id_001",
    }
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    row = table_rows(mock_db_connection, "observations")[0]
    assert row["source_id"] == "external-id"
    assert row["source_id_001"] == "source-value"

    merged = stage_import(
        mock_db_connection,
        upload("rows.csv", "_id,source_id\nexternal-id,updated\n"),
        "merge",
        "observations",
    )
    assert merged["source_mapping"] == staged["source_mapping"]
    assert merged["eligible_identity_fields"] == ["_id", "source_id"]
    review = preview_import(mock_db_connection, merged["import_id"], ["_id"])
    assert review["updated"] == 1 and review["added"] == 0
    confirm(mock_db_connection, merged, review)
    result = table_rows(mock_db_connection, "observations")[0]
    assert result["_id"] == row["_id"]
    assert result["source_id_001"] == "updated"


def test_append_maps_legacy_name_to_existing_physical_column(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, "Bird_Name" TEXT)'
        )

    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "Bird Name\nHeron\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    row = table_rows(mock_db_connection, "observations")[0]
    assert row["Bird_Name"] == "Heron"


@pytest.mark.parametrize("source", ["locusmap_favorites.gpx", "my_shapefile_data.zip"])
def test_zip_ignores_system_metadata(source):
    path = Path(__file__).parent / "assets" / source
    if path.suffix == ".zip":
        original = path.read_bytes()
    else:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("_sample.gpx", path.read_bytes())
        original = buffer.getvalue()
    expected = importer._parse_zip(original)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("__MACOSX/nested/metadata", "not data")
        archive.writestr("nested/._sample.gpx", "not GPX")
        archive.writestr("nested/.DS_Store", "not data")
        with zipfile.ZipFile(io.BytesIO(original)) as source_archive:
            for member in source_archive.infolist():
                archive.writestr(member, source_archive.read(member))
    assert importer._parse_zip(buffer.getvalue()) == expected


def test_zip_rejects_metadata_only_archive():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("__MACOSX/._sample.gpx", "metadata")
    with pytest.raises(ImportValidationError, match="no data files"):
        importer._parse_zip(buffer.getvalue())


def test_zip_validates_metadata_paths_before_filtering():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("__MACOSX/../._sample.gpx", "metadata")
        archive.writestr("rows.csv", "name\nHeron\n")
    with pytest.raises(ImportValidationError, match="unsafe"):
        importer._parse_zip(buffer.getvalue())


def test_zip_counts_metadata_toward_expanded_size_limit(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(".DS_Store", "metadata")
        archive.writestr("rows.csv", "name\nHeron\n")
    monkeypatch.setattr(importer, "MAX_EXPANDED_BYTES", len("name\nHeron\n"))
    with pytest.raises(ImportValidationError, match="100 MiB"):
        importer._parse_zip(buffer.getvalue())


def test_zip_rejects_unsupported_members_without_staging(mock_db_connection):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("rows.csv", "name\nHeron\n")
        archive.writestr("notes.txt", "unsupported")

    with pytest.raises(ImportValidationError, match="unsupported"):
        stage_import(
            mock_db_connection,
            upload_bytes("mixed.zip", buffer.getvalue()),
            "create",
            "observations",
        )
    assert check_dataset_name(mock_db_connection, "observations")["available"] is True


def test_zip_rejects_unsafe_paths(mock_db_connection):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("../rows.csv", "name\nHeron\n")

    with pytest.raises(ImportValidationError, match="unsafe"):
        stage_import(
            mock_db_connection,
            upload_bytes("unsafe.zip", buffer.getvalue()),
            "create",
            "observations",
        )


def test_multi_file_zip_appends_rows_in_archive_order(mock_db_connection):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("first.csv", "name\nHeron\n")
        archive.writestr("second.json", '[{"name":"Ibis"}]')

    staged = stage_import(
        mock_db_connection,
        upload_bytes("birds.zip", buffer.getvalue()),
        "create",
        "observations",
    )
    assert staged["source_format"] == "zip"
    assert staged["record_count"] == 2
    preview = preview_import(mock_db_connection, staged["import_id"])
    confirm(mock_db_connection, staged, preview)

    expected_ids = [
        hashlib.md5(f"{staged['import_id']}:{ordinal}".encode()).hexdigest()
        for ordinal in (1, 2)
    ]
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute('SELECT _id, name FROM "public"."observations"')
        rows_by_id = dict(cursor.fetchall())
    assert rows_by_id[expected_ids[0]] == "Heron"
    assert rows_by_id[expected_ids[1]] == "Ibis"


def test_zip_infers_coordinates_for_each_tabular_member(mock_db_connection):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "first.csv", "name,longitude,latitude\nHeron,-54.1,3.2\n"
        )
        archive.writestr(
            "second.csv",
            "name,decimalLongitude,decimalLatitude\nIbis,-54.2,3.3\n",
        )

    staged = stage_import(
        mock_db_connection,
        upload_bytes("points.zip", buffer.getvalue()),
        "create",
        "zip_points",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])

    assert preview["geometry_valid"] == 2
    assert preview["geometry_invalid"] == 0


def test_successful_import_archives_exact_source(
    mock_db_connection, importer_datalake
):
    source = b"name\nHeron\n"
    staged = stage_import(
        mock_db_connection,
        upload_bytes("birds.csv", source),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    result = confirm(mock_db_connection, staged, preview)

    archived = Path(result["archive_path"])
    assert archived.is_relative_to(importer_datalake)
    assert archived.read_bytes() == source
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT status, source_data FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        assert cursor.fetchone() == ("archived", None)


def test_archive_failure_does_not_write_target_and_can_be_retried(
    mock_db_connection, importer_datalake, tmp_path, monkeypatch
):
    blocked_root = tmp_path / "blocked"
    blocked_root.write_text("not a directory")
    monkeypatch.setenv("DATASET_IMPORTER_DATALAKE_ROOT", str(blocked_root))
    staged = stage_import(
        mock_db_connection,
        upload("birds.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    with pytest.raises(OSError):
        confirm(mock_db_connection, staged, preview)

    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("SELECT to_regclass('public.observations')")
        assert cursor.fetchone() == (None,)
        cursor.execute(
            "SELECT status, source_data IS NOT NULL FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        assert cursor.fetchone() == ("reviewed", True)

    monkeypatch.setenv("DATASET_IMPORTER_DATALAKE_ROOT", str(importer_datalake))
    result = confirm(mock_db_connection, staged, preview)
    assert Path(result["archive_path"]).read_bytes() == b"name\nHeron\n"


def test_database_failure_removes_archive_and_rolls_back_target(
    mock_db_connection, importer_datalake, monkeypatch
):
    staged = stage_import(
        mock_db_connection,
        upload("birds.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])

    def fail_mapping(*_args):
        raise RuntimeError("database failure after archival")

    monkeypatch.setattr(importer, "_ensure_dataset_mapping", fail_mapping)
    with pytest.raises(RuntimeError, match="database failure after archival"):
        confirm(mock_db_connection, staged, preview)

    assert list(importer_datalake.rglob("*birds.csv")) == []
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("SELECT to_regclass('public.observations')")
        assert cursor.fetchone() == (None,)
        cursor.execute(
            "SELECT status, source_data IS NOT NULL FROM dataset_importer_v2.import_sessions WHERE import_id = %s",
            (staged["import_id"],),
        )
        assert cursor.fetchone() == ("reviewed", True)


def test_geojson_reserved_properties_are_rejected(mock_db_connection):
    source = """{
        "type":"FeatureCollection",
        "features":[{
            "type":"Feature",
            "properties":{"g__type":"user value"},
            "geometry":{"type":"Point","coordinates":[1,2]}
        }]
    }"""
    with pytest.raises(ImportValidationError, match="reserved fields"):
        stage_import(
            mock_db_connection,
            upload("points.geojson", source),
            "create",
            "observations",
        )


def test_source_size_is_checked_before_base64_decode(
    mock_db_connection, monkeypatch
):
    payload = upload_bytes("rows.csv", b"four")
    monkeypatch.setattr(importer, "MAX_SOURCE_BYTES", 3)

    def fail_decode(*args, **kwargs):
        raise AssertionError("oversized payload was decoded")

    monkeypatch.setattr(importer.base64, "b64decode", fail_decode)
    with pytest.raises(ImportValidationError, match="25 MiB"):
        stage_import(mock_db_connection, payload, "create", "observations")


def test_zip_expanded_size_limit_is_enforced(mock_db_connection, monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("rows.csv", "name\nHeron\n")
    monkeypatch.setattr(importer, "MAX_EXPANDED_BYTES", 3)
    with pytest.raises(ImportValidationError, match="100 MiB"):
        stage_import(
            mock_db_connection,
            upload_bytes("rows.zip", buffer.getvalue()),
            "create",
            "observations",
        )


def test_final_column_limit_includes_internal_id(mock_db_connection):
    headers = ",".join(f"column_{index}" for index in range(150))
    values = ",".join("value" for _ in range(150))
    with pytest.raises(ImportValidationError, match="at most 150 columns"):
        stage_import(
            mock_db_connection,
            upload("wide.csv", f"{headers}\n{values}\n"),
            "create",
            "wide_dataset",
        )


def test_repreview_invalidates_the_previous_confirmation(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    first = preview_import(mock_db_connection, staged["import_id"])
    second = preview_import(mock_db_connection, staged["import_id"])
    assert first["preview_id"] != second["preview_id"]

    with pytest.raises(ImportValidationError, match="no longer current"):
        confirm(mock_db_connection, staged, first)
    assert check_dataset_name(mock_db_connection, "observations")["available"] is False
    assert confirm(mock_db_connection, staged, second)["success"] is True


def test_new_column_on_matching_row_is_counted_as_update(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s)', ("one", "A")
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,note\nA,new\n"),
        "merge",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert preview["updated"] == 1
    assert preview["columns_added"] == 1
    confirm(mock_db_connection, staged, preview)
    assert table_rows(mock_db_connection, "observations")[0]["note"] == "new"


def test_unchanged_match_is_not_physically_updated(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT, note TEXT)'
        )
        cursor.execute('CREATE TABLE "public"."update_audit" (count INTEGER)')
        cursor.execute('INSERT INTO "public"."update_audit" VALUES (0)')
        cursor.execute(
            """CREATE FUNCTION count_observation_updates() RETURNS trigger AS $$
            BEGIN UPDATE public.update_audit SET count = count + 1; RETURN NEW; END;
            $$ LANGUAGE plpgsql"""
        )
        cursor.execute(
            """CREATE TRIGGER observations_updated AFTER UPDATE ON public.observations
            FOR EACH ROW EXECUTE FUNCTION count_observation_updates()"""
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s, %s)',
            ("one", "A", "same"),
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,note\nA,same\n"),
        "merge",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert preview["updated"] == 0
    confirm(mock_db_connection, staged, preview)
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute("SELECT count FROM public.update_audit")
        assert cursor.fetchone()[0] == 0


@pytest.mark.parametrize("goal", ["append", "merge"])
def test_empty_append_and_merge_are_noops_without_identity(
    mock_db_connection, goal
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
        cursor.execute(
            'INSERT INTO "public"."observations" VALUES (%s, %s)', ("one", "A")
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "new_field\n"),
        goal,
        "observations",
    )
    assert staged["eligible_identity_fields"] == []
    preview = preview_import(mock_db_connection, staged["import_id"])
    assert preview["added"] == preview["updated"] == preview["deleted"] == 0
    assert preview["unchanged"] == preview["final_count"] == 1
    confirm(mock_db_connection, staged, preview)
    assert [(row["_id"], row["code"]) for row in table_rows(mock_db_connection, "observations")] == [("one", "A")]
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT 1 FROM information_schema.columns WHERE table_schema = 'public' AND table_name = 'observations' AND column_name = 'new_field'"
        )
        assert cursor.fetchone() is None


@pytest.mark.parametrize("goal", ["create", "sync"])
def test_empty_create_and_sync_are_rejected(mock_db_connection, goal):
    target = "new_observations"
    if goal == "sync":
        target = "observations"
        with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
            cursor.execute(
                'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
            )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code\n"),
        goal,
        target,
    )

    with pytest.raises(ImportValidationError, match="cannot be empty"):
        preview_import(mock_db_connection, staged["import_id"])


def test_existing_dataset_final_column_limit_is_enforced(
    mock_db_connection, monkeypatch
):
    monkeypatch.setattr(importer, "MAX_COLUMNS", 3)
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, existing TEXT)'
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "note,region\nnew,north\n"),
        "append",
        "observations",
    )

    with pytest.raises(ImportValidationError, match="exceed 150 columns"):
        preview_import(mock_db_connection, staged["import_id"])


def test_schema_change_after_review_rejects_confirmation(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'CREATE TABLE "public"."observations" (_id TEXT PRIMARY KEY, code TEXT)'
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code\nA\n"),
        "append",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            'ALTER TABLE "public"."observations" ALTER COLUMN code SET DEFAULT \'unknown\''
        )
    with pytest.raises(ImportValidationError, match="changed after review"):
        confirm(mock_db_connection, staged, preview)


def test_completed_confirmation_is_idempotent(mock_db_connection):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    first = confirm(mock_db_connection, staged, preview)
    second = confirm(mock_db_connection, staged, preview)

    assert second == first
    assert len(table_rows(mock_db_connection, "observations")) == 1


def test_concurrent_confirmation_is_idempotent(
    mock_db_connection, monkeypatch
):
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "name\nHeron\n"),
        "create",
        "observations",
    )
    preview = preview_import(mock_db_connection, staged["import_id"])
    initial_reads = threading.Barrier(2)
    original_status = importer._import_status

    def synchronize_status_read(*args):
        status = original_status(*args)
        initial_reads.wait(timeout=5)
        return status

    monkeypatch.setattr(importer, "_import_status", synchronize_status_read)
    with ThreadPoolExecutor(max_workers=2) as executor:
        first = executor.submit(confirm, mock_db_connection, staged, preview)
        second = executor.submit(confirm, mock_db_connection, staged, preview)
        results = [first.result(timeout=5), second.result(timeout=5)]

    assert results[0] == results[1]
    assert len(table_rows(mock_db_connection, "observations")) == 1


def identity_target(db, records):
    with psycopg.connect(db) as conn, conn.cursor() as cursor:
        cursor.execute(
            "CREATE TABLE public.identities (_id TEXT PRIMARY KEY, code TEXT, part TEXT, note TEXT)"
        )
        cursor.executemany(
            "INSERT INTO public.identities VALUES (%s, %s, %s, %s)", records
        )


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize(
    "filename,contents,fields,ordinal",
    [
        ("rows.json", '[{"code":"ok"},{"note":"private"}]', ["code"], 2),
        ("rows.json", '[{"code":null,"note":"private"}]', ["code"], 1),
        ("rows.json", '[{"code":"","note":"private"}]', ["code"], 1),
        ("rows.json", '[{"code":" \\t\\n","note":"private"}]', ["code"], 1),
        ("rows.csv", "code,note\n,private\n", ["code"], 1),
        (
            "rows.json",
            json.dumps([{"code": "\u00a0\u202f\u0085\u2003"}]),
            ["code"],
            1,
        ),
        ("rows.json", '[{"code":"ok","part":null}]', ["code", "part"], 1),
    ],
)
def test_incomplete_upload_rejected_without_target_changes(
    mock_db_connection, goal, filename, contents, fields, ordinal
):
    identity_target(mock_db_connection, [("internal", None, None, "unchanged")])
    before = table_rows(mock_db_connection, "identities")
    staged = stage_import(
        mock_db_connection, upload(filename, contents), goal, "identities"
    )
    with pytest.raises(
        ImportValidationError, match="Incomplete uploaded identity"
    ) as error:
        preview_import(mock_db_connection, staged["import_id"], fields)
    assert f"selected field '{fields[-1]}'" in str(error.value)
    assert f"ordinal(s): {ordinal}." in str(error.value)
    assert "private" not in str(error.value)
    assert table_rows(mock_db_connection, "identities") == before


def test_incomplete_error_uses_original_field_and_bounded_ordinals(mock_db_connection):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "CREATE TABLE public.identities (_id TEXT PRIMARY KEY, source_id TEXT)"
        )
    staged = stage_import(
        mock_db_connection,
        upload("rows.json", json.dumps([{"_id": None, "note": "private"}] * 20)),
        "merge",
        "identities",
    )
    with pytest.raises(ImportValidationError) as error:
        preview_import(mock_db_connection, staged["import_id"], ["_id"])
    assert "selected field '_id'" in str(error.value)
    assert "1, 2, 3, 4, 5, 6, 7, 8, 9, 10, …" in str(error.value)
    assert "11" not in str(error.value) and "private" not in str(error.value)


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
def test_populated_identity_values_use_exact_stored_text(
    mock_db_connection, goal, policy
):
    values = [0, False, "null", "undefined", " A ", "A", "a"]
    stored = ["0", "false", "null", "undefined", " A ", "A", "a"]
    identity_target(
        mock_db_connection,
        [(str(i), value, None, "old") for i, value in enumerate(stored)],
    )
    staged = stage_import(
        mock_db_connection,
        upload(
            "rows.json",
            json.dumps([{"code": value, "note": ""} for value in values]),
        ),
        goal,
        "identities",
    )
    review = preview_import(mock_db_connection, staged["import_id"], ["code"], policy)
    assert review["added"] == review["deleted"] == 0
    assert review["updated"] == (len(values) if policy == "imported" else 0)
    confirm(mock_db_connection, staged, review)
    assert {row["code"] for row in table_rows(mock_db_connection, "identities")} == set(
        stored
    )

    assert {
        row["note"] for row in table_rows(mock_db_connection, "identities")
    } == ({""} if policy == "imported" else {"old"})


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
@pytest.mark.parametrize("composite", [False, True])
def test_matching_target_duplicates_rejected_even_when_identical(
    mock_db_connection, goal, policy, composite
):
    identity_target(
        mock_db_connection, [("one", "A", "X", "same"), ("two", "A", "X", "same")]
    )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,part,note\nA,X,same\n"),
        goal,
        "identities",
    )
    before = table_rows(mock_db_connection, "identities")
    with pytest.raises(
        ImportValidationError,
        match="Multiple existing records.*match an uploaded identity",
    ):
        preview_import(
            mock_db_connection,
            staged["import_id"],
            ["code", "part"] if composite else ["code"],
            policy,
        )
    assert table_rows(mock_db_connection, "identities") == before


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
def test_unmatched_duplicates_and_incomplete_targets_follow_goal(
    mock_db_connection, goal, policy
):
    records = [
        ("matched", "A", "X", "old"),
        ("dup1", "B", "X", "old"),
        ("dup2", "B", "X", "old"),
    ]
    records += [
        (f"blank{i}", value, "X", "old")
        for i, value in enumerate([None, None, "", "", " \t\n", " \t\n", "\u00a0", "\u202f"])
    ]
    # One missing component makes even an otherwise populated target identity incomplete.
    records += [("missing-part", "A", None, "old")]
    identity_target(mock_db_connection, records)
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,part,note\nA,X,new\nC,X,added\n"),
        goal,
        "identities",
    )
    review = preview_import(
        mock_db_connection, staged["import_id"], ["code", "part"], policy
    )
    assert review["deleted"] == (len(records) - 1 if goal == "sync" else 0)
    assert review["added"] == 1
    assert review["updated"] == (1 if policy == "imported" else 0)
    confirm(mock_db_connection, staged, review)
    rows = table_rows(mock_db_connection, "identities")
    assert (
        len(rows)
        == review["final_count"]
        == (2 if goal == "sync" else len(records) + 1)
    )
    assert next(row for row in rows if row["_id"] == "matched")["note"] == (
        "new" if policy == "imported" else "old"
    )


@pytest.mark.parametrize("goal", ["merge", "sync"])
def test_unique_combinations_allow_repeated_components(mock_db_connection, goal):
    identity_target(
        mock_db_connection,
        [
            ("one", "A", "X", "old"),
            ("two", "A", "Y", "old"),
            ("three", "B", "X", "old"),
        ],
    )
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,part,note\nA,X,new\nA,Y,new\nB,X,new\n"),
        goal,
        "identities",
    )
    review = preview_import(mock_db_connection, staged["import_id"], ["code", "part"])
    assert review["updated"] == 3 and review["added"] == review["deleted"] == 0
    confirm(mock_db_connection, staged, review)


@pytest.mark.parametrize("goal", ["merge", "sync"])
def test_duplicate_uploaded_combinations_rejected(mock_db_connection, goal):
    identity_target(mock_db_connection, [])
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,part\nA,X\nA,X\n"),
        goal,
        "identities",
    )
    with pytest.raises(
        ImportValidationError, match="Uploaded identity combinations must be unique"
    ):
        preview_import(mock_db_connection, staged["import_id"], ["code", "part"])
    assert table_rows(mock_db_connection, "identities") == []


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
def test_no_matches_adds_upload_and_keeps_or_deletes_target(
    mock_db_connection, goal, policy
):
    identity_target(
        mock_db_connection, [("one", "B", None, "old"), ("two", "B", None, "old")]
    )
    staged = stage_import(
        mock_db_connection, upload("rows.csv", "code,note\nA,new\n"), goal, "identities"
    )
    review = preview_import(mock_db_connection, staged["import_id"], ["code"], policy)
    assert review["added"] == 1 and review["updated"] == 0
    assert review["deleted"] == (2 if goal == "sync" else 0)
    confirm(mock_db_connection, staged, review)
    assert (
        len(table_rows(mock_db_connection, "identities"))
        == review["final_count"]
        == (1 if goal == "sync" else 3)
    )


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("policy", ["imported", "existing"])
@pytest.mark.parametrize("field,stored", [("_id", "source_id"), ("_uuid", "_uuid")])
def test_sixteen_target_rows_three_uploaded_kobo_records(
    mock_db_connection, goal, policy, field, stored
):
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            psycopg.sql.SQL(
                "CREATE TABLE public.kobo (_id TEXT PRIMARY KEY, {} TEXT, note TEXT)"
            ).format(psycopg.sql.Identifier(stored))
        )
        cursor.executemany(
            "INSERT INTO public.kobo VALUES (%s, %s, %s)",
            [(f"internal-{i}", str(i) if i < 3 else None, "old") for i in range(16)],
        )
    contents = json.dumps([{field: str(i), "note": "new"} for i in range(3)])
    for iteration in range(2):
        staged = stage_import(
            mock_db_connection, upload("kobo.json", contents), goal, "kobo"
        )
        assert staged["source_mapping"][field] == stored
        review = preview_import(
            mock_db_connection, staged["import_id"], [field], policy
        )
        assert review["deleted"] == (13 if goal == "sync" and iteration == 0 else 0)
        assert review["updated"] == (
            3 if policy == "imported" and iteration == 0 else 0
        )
        assert review["added"] == review["columns_added"] == 0
        assert review["unchanged"] == (16 if goal == "merge" else 3) - review[
            "updated"
        ]
        first = confirm(mock_db_connection, staged, review)
        assert confirm(mock_db_connection, staged, review) == first
        rows = table_rows(mock_db_connection, "kobo")
        assert len(rows) == review["final_count"] == (16 if goal == "merge" else 3)
        assert len({row["_id"] for row in rows}) == len(rows)
        for i in range(3):
            row = next(row for row in rows if row[stored] == str(i))
            assert row["_id"] == f"internal-{i}"
            assert row["note"] == ("new" if policy == "imported" else "old")


@pytest.mark.parametrize("goal", ["merge", "sync"])
@pytest.mark.parametrize("marker", [None, 1])
def test_pending_old_rule_preview_requires_review_again(
    mock_db_connection, goal, marker, importer_datalake
):
    identity_target(mock_db_connection, [("one", "A", None, "old")])
    staged = stage_import(
        mock_db_connection, upload("rows.csv", "code,note\nA,new\n"), goal, "identities"
    )
    review = preview_import(mock_db_connection, staged["import_id"], ["code"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        old = dict(review)
        old.pop("identity_rules_version")
        if marker is not None:
            old["identity_rules_version"] = marker
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET preview = %s WHERE import_id = %s",
            (json.dumps(old), staged["import_id"]),
        )
    before = table_rows(mock_db_connection, "identities")
    with pytest.raises(
        ImportValidationError, match="Identity rules changed.*Review the import again"
    ):
        confirm(mock_db_connection, staged, review)
    assert table_rows(mock_db_connection, "identities") == before
    assert not list(importer_datalake.rglob("*.csv"))
    current = preview_import(mock_db_connection, staged["import_id"], ["code"])
    assert current["identity_rules_version"] == importer.IDENTITY_RULES_VERSION
    assert confirm(mock_db_connection, staged, current)["success"]
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        cursor.execute(
            "UPDATE dataset_importer_v2.import_sessions SET preview = preview - 'identity_rules_version' WHERE import_id = %s",
            (staged["import_id"],),
        )
    assert confirm(mock_db_connection, staged, current)["success"]
    assert len(table_rows(mock_db_connection, "identities")) == 1


@pytest.mark.parametrize("mutation", ["incomplete", "duplicate"])
def test_apply_revalidates_staged_identities_under_target_lock(
    mock_db_connection, mutation
):
    identity_target(mock_db_connection, [("one", "A", None, "old")])
    staged = stage_import(
        mock_db_connection,
        upload("rows.csv", "code,note\nA,new\n"),
        "sync",
        "identities",
    )
    review = preview_import(mock_db_connection, staged["import_id"], ["code"])
    with psycopg.connect(mock_db_connection) as conn, conn.cursor() as cursor:
        if mutation == "incomplete":
            cursor.execute(
                "UPDATE dataset_importer_v2.import_rows SET payload = payload - 'code' WHERE import_id = %s",
                (staged["import_id"],),
            )
        else:
            cursor.execute(
                "INSERT INTO dataset_importer_v2.import_rows SELECT import_id, 2, payload FROM dataset_importer_v2.import_rows WHERE import_id = %s",
                (staged["import_id"],),
            )
    before = table_rows(mock_db_connection, "identities")
    with pytest.raises(
        ImportValidationError,
        match="Incomplete uploaded identity|Uploaded identity combinations must be unique",
    ):
        confirm(mock_db_connection, staged, review)
    assert table_rows(mock_db_connection, "identities") == before


@pytest.mark.parametrize("goal", ["create", "append"])
def test_create_append_ignore_incomplete_and_duplicate_identities(
    mock_db_connection, goal
):
    if goal == "append":
        identity_target(mock_db_connection, [])
    staged = stage_import(
        mock_db_connection,
        upload("rows.json", '[{"code":null},{"code":null},{"code":"A"},{"code":"A"}]'),
        goal,
        "identities",
    )
    review = preview_import(mock_db_connection, staged["import_id"])
    assert review["added"] == 4
    confirm(mock_db_connection, staged, review)
    assert len(table_rows(mock_db_connection, "identities")) == 4
